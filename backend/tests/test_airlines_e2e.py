"""End-to-end test of the Airlines page's engine entry point: synthetic files
-> detection -> full run -> known matched / unmatched / refund numbers."""
import io
import os

import openpyxl
import pandas as pd
import pytest

from backend.airlines import engine
from backend.tests.airlines_fixtures import build_sample_files


@pytest.fixture()
def sample(tmp_path, monkeypatch):
    # isolate the carry-forward state so tests never touch real state
    monkeypatch.setattr(engine, "STATE_DIR", str(tmp_path / "state"))

    # config.yaml uses Windows drive paths (C:/ReconState/...). On Windows
    # those are used verbatim, bypassing STATE_DIR, so force bare file names.
    _orig_resolve = engine.resolve_state_path

    def _isolated_resolve(state_file):
        if not state_file:
            return None
        name = os.path.basename(str(state_file).replace("\\", "/"))
        return _orig_resolve(name)

    monkeypatch.setattr(engine, "resolve_state_path", _isolated_resolve)

    folder = tmp_path / "in"
    paths = build_sample_files(str(folder))
    uploads = []
    for p in paths.values():
        with open(p, "rb") as fh:
            uploads.append((os.path.basename(p), fh.read()))
    return uploads


def _by_airline(results):
    return {r["airline_key"]: r for r in results}


def test_detection_ignores_file_names(sample):
    renamed = [(f"random_{i}.xlsx", data) for i, (_, data) in enumerate(sample)]
    detected = engine.detect_uploaded_files(renamed)
    assert set(detected) == {"ledger_a", "ledger_b", "airline_1", "airline_2",
                             "airline_3", "cancellation_report", "fonepay_refund"}


def test_missing_required_file_is_a_readable_error(sample):
    without_shree = [u for u in sample if u[0] != "shree.xlsx"]
    with pytest.raises(RuntimeError, match="airline_1"):
        engine.detect_uploaded_files(without_shree)


def test_full_run_known_numbers(sample):
    results, detected, log = engine.run_airline_reconciliation(sample)
    r = _by_airline(results)
    assert len(results) == 3

    shree = r["airline_1"]
    assert list(shree["recon"]) == ["eSewa vs Shree Airlines", "Fonepay vs Shree Airlines"]
    assert len(shree["recon"]["eSewa vs Shree Airlines"]) == 2       # S1, S3
    assert len(shree["recon"]["Fonepay vs Shree Airlines"]) == 1     # S2
    un = shree["unrecon"]["Unrecon"]
    assert sorted(un["Ticket"]) == ["1110000000004", "1110000000005"]
    assert dict(zip(un["Ticket"], un["Reverted?"])) == {"1110000000004": "Yes", "1110000000005": "No"}
    assert shree["summary"]["Summary eSewa"].iloc[0]["Difference"] == 50

    yeti = r["airline_2"]
    assert yeti["summary"]["Summary eSewa"].iloc[0]["Total Airline Tickets"] == 3   # duplicate removed
    assert list(yeti["unrecon"]["Unrecon"]["Ticket"]) == ["2220000000003"]

    buddha = r["airline_3"]
    assert "eSewa Refund Recon" in buddha["recon"]
    assert "Fonepay Refund Recon" in buddha["recon"]
    assert buddha["recon"]["eSewa Refund Recon"]["Cashback Per Ticket"].iloc[0] == 100
    refund_un = buddha["unrecon"]["Refund Unrecon"]
    assert list(refund_un["TICKET NO"]) == ["3330000000006"]
    assert list(refund_un["Refund Status"]) == ["Overdue"]
    assert list(buddha["unrecon"]["Unrecon"]["Refund?"]) == ["No"]


def test_reports_are_valid_excel(sample):
    results, _, _ = engine.run_airline_reconciliation(sample)
    for r in results:
        wb = openpyxl.load_workbook(io.BytesIO(r["bytes"]))
        assert "Unrecon" in wb.sheetnames
        assert wb.sheetnames[0].startswith("Summary")


def test_pending_refund_carries_forward_between_runs(sample):
    engine.run_airline_reconciliation(sample)
    state = os.path.join(engine.STATE_DIR, "pending_refunds.csv")
    assert os.path.exists(state)
    assert list(pd.read_csv(state, dtype=str)["TICKET NO"]) == ["3330000000006"]

    # second run (same files) must not duplicate the pending refund
    engine.run_airline_reconciliation(sample)
    assert list(pd.read_csv(state, dtype=str)["TICKET NO"]) == ["3330000000006"]


def test_config_edit_changes_names(sample):
    cfg = engine.load_config()
    cfg["names"]["airline_1"] = "Renamed Air"
    results, _, _ = engine.run_airline_reconciliation(sample, cfg=cfg)
    assert _by_airline(results)["airline_1"]["airline_name"] == "Renamed Air"


# ---- HTML-table ".xls" exports (some airline systems save HTML as .xls) ----

def _yeti_as_html_xls(encoding):
    rows = [["Yeti sales"],
            ["Issue", "From", "To", "Flight", "Date", "Ticket", "Passenger", "Fare Code", "Net Total"],
            ["2026-09-30", "KTM", "BIR", "YT-2", "2026-09-30", "2220000000001", "Pax", "Y", "1200"],
            ["2026-09-30", "KTM", "BIR", "YT-2", "2026-09-30", "2220000000002", "Pax", "Y", "1,300"],
            ["2026-09-30", "KTM", "BIR", "YT-2", "2026-09-30", "2220000000002", "Pax", "Y", "1,300"],
            ["2026-09-30", "KTM", "BIR", "YT-2", "2026-09-30", "2220000000003", "Pax", "Y", "999"]]
    body = "".join(
        "<tr>" + ("<td colspan='9'>%s</td>" % r[0] if len(r) == 1 else "".join(f"<td>{c}</td>" for c in r)) + "</tr>"
        for r in rows)
    html = f"<html><body><table>{body}</table></body></html>"
    return html.encode(encoding)


@pytest.mark.parametrize("encoding", ["utf-8", "utf-16"])
@pytest.mark.parametrize("lxml_available", [True, False])
def test_html_table_xls_loads(sample, monkeypatch, encoding, lxml_available):
    if not lxml_available:
        def _no_lxml(*a, **k):
            raise ImportError("Import lxml failed. Use pip or conda to install the lxml package.")
        monkeypatch.setattr(pd, "read_html", _no_lxml)
    uploads = [(n, d) for n, d in sample if n != "yeti.xlsx"]
    uploads.append(("Yeti.xls", _yeti_as_html_xls(encoding)))

    detected = engine.detect_uploaded_files(uploads)
    assert detected["airline_2"] == "Yeti.xls"

    results, _, _ = engine.run_airline_reconciliation(uploads)
    yeti = _by_airline(results)["airline_2"]
    assert yeti["summary"]["Summary eSewa"].iloc[0]["Total Airline Tickets"] == 3
    assert len(yeti["recon"]["eSewa vs Yeti Airlines"]) == 1
    assert len(yeti["recon"]["Fonepay vs Yeti Airlines"]) == 1
    assert yeti["summary"]["Summary Fonepay"].iloc[0]["Airline Amount"] == 1300   # "1,300" parsed
    assert list(yeti["unrecon"]["Unrecon"]["Ticket"]) == ["2220000000003"]
