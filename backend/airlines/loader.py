"""
loader.py
Reads a source file into a clean pandas DataFrame:
- unmerges merged cells (fills value across the original merge range)
- skips to the configured header row(s)
- for 2-row headers (airline_3), combines them into single column names
- drops fully blank rows
- drops subtotal/total rows (blank match-key OR a "Total"/"Subtotal" label)
"""
import os
import re
from html.parser import HTMLParser

import pandas as pd
import openpyxl


def _missing_xlrd_error(path):
    return RuntimeError(
        f"Could not read legacy Excel file {path}. Install xlrd with "
        "'python -m pip install xlrd', then run again."
    )


def _unmerge_and_fill(ws):
    """Fill merged cell ranges with their top-left value so every cell
    in the range reads correctly when converted to rows."""
    merged_ranges = list(ws.merged_cells.ranges)
    for merged_range in merged_ranges:
        min_col, min_row, max_col, max_row = merged_range.bounds
        top_left_value = ws.cell(row=min_row, column=min_col).value
        ws.unmerge_cells(str(merged_range))
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                ws.cell(row=r, column=c, value=top_left_value)


def _load_xlsx_raw(path, sheet_name=None):
    wb = openpyxl.load_workbook(path, data_only=True)
    if sheet_name:
        if sheet_name not in wb.sheetnames:
            raise ValueError(
                f"Sheet '{sheet_name}' not found in {path}. Available sheets: {wb.sheetnames}"
            )
        ws = wb[sheet_name]
    else:
        ws = wb[wb.sheetnames[0]]
    _unmerge_and_fill(ws)
    data = list(ws.iter_rows(values_only=True))
    return data


def _load_xls_raw(path):
    try:
        import xlrd
    except ImportError as exc:
        raise _missing_xlrd_error(path) from exc

    try:
        book = xlrd.open_workbook(path, formatting_info=True)
    except xlrd.biffh.XLRDError:
        return _load_html_table_raw(path)

    sheet = book.sheet_by_index(0)
    data = [sheet.row_values(r) for r in range(sheet.nrows)]

    for row_start, row_end, col_start, col_end in sheet.merged_cells:
        value = data[row_start][col_start]
        for r in range(row_start, row_end):
            for c in range(col_start, col_end):
                data[r][c] = value

    return data


class _TableParser(HTMLParser):
    """Minimal stdlib HTML table reader (no lxml / bs4 needed). Collects
    every <table> as a list of rows; colspan is expanded by repeating the
    cell value (same effect as the merged-cell fill used for .xlsx)."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tables, self._stack = [], []
        self._row = self._cell = None
        self._span = 1

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self._stack.append([])
        elif tag == "tr" and self._stack:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell, self._span = [], 1
            try:
                self._span = max(1, int(dict(attrs).get("colspan") or 1))
            except ValueError:
                pass
        elif tag == "br" and self._cell is not None:
            self._cell.append(" ")

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._cell is not None and self._row is not None:
            text = "".join(self._cell).strip()
            self._row.extend([text if text else None] * self._span)
            self._cell = None
        elif tag == "tr" and self._row is not None and self._stack:
            self._stack[-1].append(self._row)
            self._row = None
        elif tag == "table" and self._stack:
            self.tables.append(self._stack.pop())


def read_html_table_rows(path):
    """Rows (list of lists) of the largest table in an HTML-as-.xls export.
    Tries pandas.read_html first (needs lxml); if that library is missing
    or fails, falls back to a built-in stdlib parser so the file still
    loads without any extra install."""
    with open(path, "rb") as f:
        raw = f.read()
    encoding = "utf-16" if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else None
    try:
        tables = pd.read_html(path, encoding=encoding, header=None)
        if tables:
            df = max(tables, key=lambda table: table.shape[0] * table.shape[1])
            df = df.astype(object).where(pd.notna(df), None)
            return df.values.tolist()
    except (ImportError, ValueError):
        pass  # no lxml/bs4, or pandas found no table -> try the built-in parser

    text = None
    for enc in ([encoding] if encoding else []) + ["utf-8-sig", "cp1252", "latin-1"]:
        try:
            text = raw.decode(enc)
            break
        except (UnicodeDecodeError, LookupError):
            continue
    parser = _TableParser()
    parser.feed(text or "")
    tables = [t for t in parser.tables if t]
    if not tables:
        raise ValueError(f"No table data found in {path}")
    rows = max(tables, key=lambda t: len(t) * max(len(r) for r in t))
    width = max(len(r) for r in rows)
    return [r + [None] * (width - len(r)) for r in rows]


def _load_html_table_raw(path):
    return read_html_table_rows(path)


def _load_csv_raw(path):
    df = pd.read_csv(path, header=None)
    return df.values.tolist()


def _build_header(raw_rows, header_rows_cfg):
    """header_rows_cfg is either an int (1-indexed single row) or a
    list of two ints (1-indexed, combined e.g. [6,7])."""
    if isinstance(header_rows_cfg, int):
        idx = header_rows_cfg - 1
        header = [str(c).strip() if c is not None else "" for c in raw_rows[idx]]
        data_start = idx + 1
    else:
        idx1, idx2 = header_rows_cfg[0] - 1, header_rows_cfg[1] - 1
        row1 = raw_rows[idx1]
        row2 = raw_rows[idx2]
        header = []
        last_group = ""
        for c1, c2 in zip(row1, row2):
            g = str(c1).strip() if c1 is not None else ""
            sub = str(c2).strip() if c2 is not None else ""
            if g:
                last_group = g
            if sub and sub.upper() != last_group.upper():
                header.append(f"{last_group}_{sub}".strip("_"))
            elif sub:
                header.append(sub)
            elif g:
                header.append(g)
            else:
                header.append("")
        data_start = idx2 + 1
    return _dedupe_headers(header), data_start


def _dedupe_headers(header):
    seen = {}
    deduped = []
    for col in header:
        base = str(col).strip() if col is not None else ""
        if not base:
            deduped.append(base)
            continue
        count = seen.get(base, 0) + 1
        seen[base] = count
        deduped.append(base if count == 1 else f"{base}_{count}")
    return deduped


def load_source(filepath, header_rows_cfg, sheet_name=None):
    """Returns a raw (unfiltered) DataFrame with proper headers applied,
    merged cells filled, and rows above the header discarded.
    sheet_name only applies to .xlsx/.xlsm files (multi-sheet workbooks);
    ignored for .xls/.csv, which are treated as single-sheet."""
    ext = os.path.splitext(filepath)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        raw_rows = _load_xlsx_raw(filepath, sheet_name=sheet_name)
    elif ext == ".xls":
        raw_rows = _load_xls_raw(filepath)
    elif ext == ".csv":
        raw_rows = _load_csv_raw(filepath)
    else:
        raise ValueError(f"Unsupported file type: {filepath}")

    header, data_start = _build_header(raw_rows, header_rows_cfg)
    data_rows = raw_rows[data_start:]
    df = pd.DataFrame(data_rows, columns=header)

    # drop unnamed/empty-header columns (common artifact of blank merged cols)
    df = df.loc[:, [c for c in df.columns if c and str(c).strip() != ""]]
    return df


def drop_junk_rows(df, match_key_col):
    """Drops fully blank rows and subtotal/total rows.
    A row is junk if the match key is blank/NaN, OR any cell in the row
    contains the word 'total'/'subtotal' as a standalone label."""
    if df.empty:
        return df

    df = df.dropna(how="all")

    # drop rows where match key is blank
    if match_key_col in df.columns:
        df = df[df[match_key_col].notna()]
        df = df[df[match_key_col].astype(str).str.strip() != ""]

    # safety net: drop rows that look like "Total"/"Subtotal" labels
    def looks_like_total_row(row):
        for v in row:
            if isinstance(v, str):
                value = v.strip().lower()
                if re.fullmatch(r"(grand\s+)?(sub)?total(\s+amount)?", value):
                    return True
                if value in ("total", "subtotal", "grand total"):
                    return True
        if match_key_col in row.index and isinstance(row[match_key_col], str):
            value = row[match_key_col].strip().lower()
            if "total" in value:
                return True
        return False

    mask = df.apply(looks_like_total_row, axis=1)
    df = df[~mask]

    return df.reset_index(drop=True)
